import pytest
import pytest_asyncio
from pathlib import Path
from httpx import AsyncClient, ASGITransport
from docx import Document
from app.main import app
from app.db.init_db import init_database
from app.db.ingest_manuals import main as ingest_main

@pytest_asyncio.fixture(autouse=True, scope="module")
async def setup_database_and_rag():
    await init_database()
    ingest_main()

def create_sample_docx(target_path: Path):
    doc = Document()
    doc.add_heading("Inspection Report", 0)
    doc.add_paragraph("Equipment: HX-204")
    doc.add_paragraph("Tube Wall Thickness (measured): 2.8 mm")
    doc.add_paragraph("Tube Wall Thickness (required min): 3.5 mm")
    doc.save(target_path)

@pytest.mark.asyncio
async def test_analyze_doc_endpoint(tmp_path):
    sample_file = tmp_path / "test_report.docx"
    create_sample_docx(sample_file)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        with open(sample_file, "rb") as f:
            files = {"file": ("test_report.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            data = {"machine_id": "HX-204", "analysis_type": "thickness_approval"}
            response = await ac.post("/api/v1/analyze-doc", files=files, data=data)

    assert response.status_code == 202
    res = response.json()
    assert res["status"] == "completed"
    assert "doc_reader" in res["tools_used"]
    assert "report_gen" in res["tools_used"]
    assert "result" in res
    assert "report_path" in res["result"]

@pytest.mark.asyncio
async def test_list_agent_runs():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        response = await ac.get("/api/v1/agent/runs")
    assert response.status_code == 200
    data = response.json()
    assert isinstance(data, list)
