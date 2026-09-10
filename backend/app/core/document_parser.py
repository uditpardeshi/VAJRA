import os
import csv
import json
from pathlib import Path
from typing import List, Dict, Any, Optional
import logging

logger = logging.getLogger(__name__)

class DocumentParser:
    """
    Universal document parser that ingests diverse formats:
    PDF, DOCX, XLSX, CSV, TXT, MD, JSON, and Images.
    Produces structured chunks with modality tags, table structures, and visual diagrams.
    """

    def __init__(self, upload_base_dir: Optional[Path] = None):
        self.upload_base_dir = upload_base_dir or (Path(__file__).parent.parent.parent / "uploads")
        self.upload_base_dir.mkdir(parents=True, exist_ok=True)

    def parse_file(
        self,
        file_path: str,
        filename: str,
        session_id: str = "default",
        chunk_size: int = 500,
        overlap: int = 50
    ) -> List[Dict[str, Any]]:
        path = Path(file_path)
        ext = path.suffix.lower()

        try:
            if ext == ".pdf":
                return self._parse_pdf(path, filename, session_id, chunk_size, overlap)
            elif ext in [".docx", ".doc"]:
                return self._parse_docx(path, filename, session_id, chunk_size, overlap)
            elif ext in [".xlsx", ".xls"]:
                return self._parse_excel(path, filename, session_id)
            elif ext == ".csv":
                return self._parse_csv(path, filename, session_id)
            elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
                return self._parse_image(path, filename, session_id)
            elif ext in [".json"]:
                return self._parse_json(path, filename, session_id, chunk_size, overlap)
            else:
                # Default text handling (.txt, .md, .log, .yaml, etc.)
                return self._parse_text(path, filename, session_id, chunk_size, overlap)
        except Exception as e:
            logger.error(f"Failed to parse {filename} ({ext}): {e}", exc_info=True)
            # Graceful fallback: treat as raw text
            return self._parse_text_fallback(path, filename, session_id)

    # ------------------ PDF PARSER ------------------
    def _parse_pdf(
        self,
        path: Path,
        filename: str,
        session_id: str,
        chunk_size: int,
        overlap: int
    ) -> List[Dict[str, Any]]:
        chunks = []
        try:
            import pymupdf as fitz
        except ImportError:
            try:
                import fitz
            except ImportError:
                fitz = None

        if fitz:
            doc = fitz.open(str(path))
            diagram_dir = self.upload_base_dir / "diagrams" / "sessions" / session_id
            diagram_dir.mkdir(parents=True, exist_ok=True)

            for page_num in range(len(doc)):
                page = doc[page_num]
                page_text = page.get_text() or ""
                p_num = page_num + 1

                # Extract images/diagrams from page if any
                for img_index, img in enumerate(page.get_images(full=True)):
                    try:
                        xref = img[0]
                        base_image = doc.extract_image(xref)
                        image_bytes = base_image["image"]
                        image_ext = base_image.get("ext", "png")
                        diag_filename = f"{path.stem}_p{p_num}_diag{img_index}.{image_ext}"
                        diag_path = diagram_dir / diag_filename

                        if not diag_path.exists():
                            with open(diag_path, "wb") as f:
                                f.write(image_bytes)

                        rel_url = f"/uploads/diagrams/sessions/{session_id}/{diag_filename}"
                        chunks.append({
                            "text": f"[DIAGRAM / SCHEMATIC extracted from {filename}, Page {p_num}]",
                            "source_file": filename,
                            "source_page": p_num,
                            "modality": "figure",
                            "image_snippet_url": rel_url,
                            "table_data": None
                        })
                    except Exception as img_err:
                        logger.debug(f"Could not extract image from PDF page: {img_err}")

                # Text chunking
                if page_text.strip():
                    for i in range(0, len(page_text), chunk_size - overlap):
                        snippet = page_text[i:i + chunk_size].strip()
                        if snippet:
                            chunks.append({
                                "text": snippet,
                                "source_file": filename,
                                "source_page": p_num,
                                "modality": "text",
                                "image_snippet_url": None,
                                "table_data": None
                            })
            doc.close()
        else:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            for p_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                for i in range(0, len(text), chunk_size - overlap):
                    snippet = text[i:i + chunk_size].strip()
                    if snippet:
                        chunks.append({
                            "text": snippet,
                            "source_file": filename,
                            "source_page": p_num,
                            "modality": "text",
                            "image_snippet_url": None,
                            "table_data": None
                        })

        return chunks

    # ------------------ DOCX PARSER ------------------
    def _parse_docx(
        self,
        path: Path,
        filename: str,
        session_id: str,
        chunk_size: int,
        overlap: int
    ) -> List[Dict[str, Any]]:
        import docx
        doc = docx.Document(str(path))
        chunks = []

        # 1. Extract paragraphs
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())

        body_text = "\n\n".join(full_text)
        for i in range(0, len(body_text), chunk_size - overlap):
            snippet = body_text[i:i + chunk_size].strip()
            if snippet:
                chunks.append({
                    "text": snippet,
                    "source_file": filename,
                    "source_page": 1,
                    "modality": "text",
                    "image_snippet_url": None,
                    "table_data": None
                })

        # 2. Extract tables
        for t_idx, table in enumerate(doc.tables, start=1):
            table_rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(cells))
            if table_rows:
                table_markdown = "\n".join(table_rows)
                chunks.append({
                    "text": f"Table {t_idx} from {filename}:\n{table_markdown}",
                    "source_file": filename,
                    "source_page": t_idx,
                    "modality": "table",
                    "image_snippet_url": None,
                    "table_data": table_markdown
                })

        return chunks

    # ------------------ EXCEL PARSER ------------------
    def _parse_excel(self, path: Path, filename: str, session_id: str) -> List[Dict[str, Any]]:
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)
        chunks = []

        for sheet_idx, sheetname in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheetname]
            rows_data = []
            for row in sheet.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    row_str = " | ".join([str(v) if v is not None else "" for v in row])
                    rows_data.append(row_str)

            if rows_data:
                # Group into batches of 25 rows for readable table context
                for b_idx in range(0, len(rows_data), 25):
                    batch = rows_data[b_idx:b_idx + 25]
                    table_str = "\n".join(batch)
                    chunks.append({
                        "text": f"Sheet '{sheetname}' ({filename}, rows {b_idx+1}-{b_idx+len(batch)}):\n{table_str}",
                        "source_file": filename,
                        "source_page": sheet_idx,
                        "modality": "table",
                        "image_snippet_url": None,
                        "table_data": table_str
                    })

        return chunks

    # ------------------ CSV PARSER ------------------
    def _parse_csv(self, path: Path, filename: str, session_id: str) -> List[Dict[str, Any]]:
        chunks = []
        rows = []
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                if any(cell.strip() for cell in row):
                    rows.append(" | ".join(row))

        if rows:
            for b_idx in range(0, len(rows), 30):
                batch = rows[b_idx:b_idx + 30]
                table_str = "\n".join(batch)
                chunks.append({
                    "text": f"CSV {filename} (rows {b_idx+1}-{b_idx+len(batch)}):\n{table_str}",
                    "source_file": filename,
                    "source_page": (b_idx // 30) + 1,
                    "modality": "table",
                    "image_snippet_url": None,
                    "table_data": table_str
                })

        return chunks

    # ------------------ IMAGE PARSER ------------------
    def _parse_image(self, path: Path, filename: str, session_id: str) -> List[Dict[str, Any]]:
        target_dir = self.upload_base_dir / "images" / "sessions" / session_id
        target_dir.mkdir(parents=True, exist_ok=True)
        dest_path = target_dir / filename

        import shutil
        if not dest_path.exists():
            shutil.copyfile(path, dest_path)

        rel_url = f"/uploads/images/sessions/{session_id}/{filename}"
        return [{
            "text": f"[IMAGE / SCHEMATIC: {filename} uploaded for session {session_id}]",
            "source_file": filename,
            "source_page": 1,
            "modality": "figure",
            "image_snippet_url": rel_url,
            "table_data": None
        }]

    # ------------------ JSON PARSER ------------------
    def _parse_json(
        self,
        path: Path,
        filename: str,
        session_id: str,
        chunk_size: int,
        overlap: int
    ) -> List[Dict[str, Any]]:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)

        formatted = json.dumps(data, indent=2)
        chunks = []
        for i in range(0, len(formatted), chunk_size - overlap):
            snippet = formatted[i:i + chunk_size].strip()
            if snippet:
                chunks.append({
                    "text": f"JSON Data from {filename}:\n{snippet}",
                    "source_file": filename,
                    "source_page": (i // chunk_size) + 1,
                    "modality": "text",
                    "image_snippet_url": None,
                    "table_data": None
                })
        return chunks

    # ------------------ TEXT PARSER ------------------
    def _parse_text(
        self,
        path: Path,
        filename: str,
        session_id: str,
        chunk_size: int,
        overlap: int
    ) -> List[Dict[str, Any]]:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()

        chunks = []
        for i in range(0, len(text), chunk_size - overlap):
            snippet = text[i:i + chunk_size].strip()
            if snippet:
                chunks.append({
                    "text": snippet,
                    "source_file": filename,
                    "source_page": (i // chunk_size) + 1,
                    "modality": "text",
                    "image_snippet_url": None,
                    "table_data": None
                })
        return chunks

    def _parse_text_fallback(self, path: Path, filename: str, session_id: str) -> List[Dict[str, Any]]:
        try:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            return [{
                "text": content[:1500],
                "source_file": filename,
                "source_page": 1,
                "modality": "text",
                "image_snippet_url": None,
                "table_data": None
            }]
        except Exception:
            return [{
                "text": f"Attached file: {filename}",
                "source_file": filename,
                "source_page": 1,
                "modality": "text",
                "image_snippet_url": None,
                "table_data": None
            }]

document_parser = DocumentParser()
