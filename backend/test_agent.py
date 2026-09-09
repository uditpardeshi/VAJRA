# test_agent.py - Run after backend starts: python test_agent.py
import httpx
import sys
from pathlib import Path

BASE = "http://localhost:8000/api/v1"

def create_test_docx():
    """Create a simple test inspection report docx."""
    from docx import Document
    doc = Document()
    doc.add_heading("Inspection Report", 0)
    doc.add_paragraph("Equipment: HX-204")
    doc.add_paragraph("Date: 2026-09-09")
    doc.add_paragraph("Inspector: Rajesh Kumar")
    doc.add_paragraph("")
    doc.add_heading("Measurements", 1)
    doc.add_paragraph("Tube Wall Thickness (measured): 2.8 mm")
    doc.add_paragraph("Tube Wall Thickness (required min): 3.5 mm")
    doc.add_paragraph("Location: Row 12, Tube 45")
    doc.add_paragraph("")
    doc.add_heading("Observations", 1)
    doc.add_paragraph("Thickness below recommended minimum. Recommend further evaluation.")
    
    upload_dir = Path("uploads")
    upload_dir.mkdir(parents=True, exist_ok=True)
    docx_path = upload_dir / "test_inspection_HX-204.docx"
    doc.save(docx_path)
    return str(docx_path)

def test_agent_analysis():
    test_file = create_test_docx()
    print(f"Created test file: {test_file}")
    
    with open(test_file, "rb") as f:
        files = {"file": (Path(test_file).name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {"machine_id": "HX-204", "analysis_type": "thickness_approval"}
        r = httpx.post(f"{BASE}/analyze-doc", files=files, data=data, timeout=120)
    
    assert r.status_code == 202, f"Status {r.status_code}: {r.text}"
    result = r.json()
    
    required = ["run_id", "status", "machine_id", "tools_used", "result", "tool_trace", "created_at"]
    for k in required:
        assert k in result, f"Missing key: {k}"
    
    assert result["status"] == "completed"
    assert "doc_reader" in result["tools_used"]
    assert "rag" in result["tools_used"]
    assert "report_gen" in result["tools_used"]
    
    res = result["result"]
    assert "recommendation" in res
    assert "report_path" in res
    assert Path(res["report_path"]).exists()
    
    print(f"✅ Agent run: {result['run_id']}")
    print(f"   Tools: {result['tools_used']}")
    print(f"   Recommendation: {res['recommendation']}")
    print(f"   Deviation: {res.get('deviation_mm')} mm")
    print(f"   Report: {res['report_path']}")
    print(f"   Trace steps: {len(result['tool_trace'])}")
    for t in result['tool_trace']:
        print(f"     - {t['tool']}: {t['duration_ms']}ms")
    
    return True

def test_agent_runs_list():
    r = httpx.get(f"{BASE}/agent/runs", timeout=10)
    assert r.status_code == 200
    runs = r.json()
    assert len(runs) > 0
    print(f"✅ Agent runs list: {len(runs)} runs")
    return True

if __name__ == "__main__":
    try:
        test_agent_analysis()
        test_agent_runs_list()
        print("\n🎉 AGENT SMOKE TESTS PASSED")
    except Exception as e:
        print(f"\n❌ FAILED: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
